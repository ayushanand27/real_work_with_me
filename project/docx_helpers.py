"""Formatting helpers that implement the course's report format (Times New Roman, 12pt body,
14pt bold left headings, 16pt bold uppercase centred title, 1.5 spacing, 1-inch margins, justified text)."""
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "Times New Roman"

def _font(run, size=12, bold=None, italic=None):
    run.font.name = FONT; run.font.size = Pt(size)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    if bold is not None: run.bold = bold
    if italic is not None: run.italic = italic

def new_document():
    d = Document()
    for s in d.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
    st = d.styles["Normal"]; st.font.name = FONT; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.line_spacing = 1.5; st.paragraph_format.space_after = Pt(0)
    return d

def center(d, text, size=12, bold=False, after=0, upper=False, italic=False):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(text.upper() if upper else text); _font(r, size, bold, italic); return p

def heading(d, text):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10); p.paragraph_format.keep_with_next = True
    _font(p.add_run(text), 14, True); return p

def subheading(d, text):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.keep_with_next = True
    _font(p.add_run(text), 12, True); return p

def para(d, text, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    """Justified body paragraph. **bold** spans are supported."""
    p = d.add_paragraph(); p.alignment = align
    for i, chunk in enumerate(text.split("**")):
        if chunk: _font(p.add_run(chunk), 12, bold=(i % 2 == 1) or None, italic=italic or None)
    return p

def bullets(d, items):
    for t in items:
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.4); p.paragraph_format.first_line_indent = Inches(-0.25)
        for i, chunk in enumerate(("•  " + t).split("**")):
            if chunk: _font(p.add_run(chunk), 12, bold=(i % 2 == 1) or None)

def numbered(d, items):
    for n, t in enumerate(items, 1):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Inches(0.4); p.paragraph_format.first_line_indent = Inches(-0.3)
        for i, chunk in enumerate((f"{n}.  " + t).split("**")):
            if chunk: _font(p.add_run(chunk), 12, bold=(i % 2 == 1) or None)

def page_break(d): d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

def _shade(cell, hexcolor):
    tcPr = cell._element.get_or_add_tcPr(); sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexcolor); tcPr.append(sh)

def table(d, header, rows, caption, widths=None, size=10, left=False):
    cp = d.add_paragraph(); cp.alignment = WD_ALIGN_PARAGRAPH.CENTER; cp.paragraph_format.keep_with_next = True
    cp.paragraph_format.line_spacing = 1.0; cp.paragraph_format.space_before = Pt(6)
    _font(cp.add_run(caption), 11, True)
    t = d.add_table(rows=1, cols=len(header)); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, h in enumerate(header):
        c = t.rows[0].cells[j]; c.text = ""; _shade(c, "D9E2F3")
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.line_spacing = 1.0
        _font(p.add_run(h), size, True)
    for row in rows:
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = ""; p = cells[j].paragraphs[0]; p.paragraph_format.line_spacing = 1.0
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if left or j == 0 or len(str(v)) > 14 else WD_ALIGN_PARAGRAPH.CENTER
            _font(p.add_run(str(v)), size)
    if widths:
        for row in t.rows:
            for j, w in enumerate(widths): row.cells[j].width = Inches(w)
    sp = d.add_paragraph(); sp.paragraph_format.line_spacing = 1.0; sp.paragraph_format.space_after = Pt(2)
    return t

def figure(d, path, caption, width=5.6):
    """Picture and caption share one paragraph (line break between) so they can never be split across pages."""
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing = 1.0; p.paragraph_format.space_after = Pt(6)
    p.add_run().add_picture(path, width=Inches(width))
    r = p.add_run(); r.add_break(); _font(r, 11)
    _font(p.add_run(caption), 11, True)

def page_numbers(d):
    for s in d.sections:
        p = s.footer.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(); _font(r, 11)
        for kind, txt in (("begin", None), (None, "PAGE"), ("end", None)):
            if kind:
                e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), kind)
            else:
                e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
            r._r.append(e)
